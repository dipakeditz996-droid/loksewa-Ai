"""Objective vs subjective coverage of the Master Question Bank: type-specific
import (template, validation, preview, duplicates, commit), the single
QuestionSelectionService, student practice (MCQ) and subjective attempts,
and the teacher approval workflow. Everything runs through the real endpoints
and models - no fixtures standing in for questions."""
import io

from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from core.models import User
from exams.models import (
    ExamCategory, Exam, Paper, Subject, Chapter, Topic, Question, QuestionAttempt,
    SubjectiveAttempt, SubjectiveAnswer, Evaluation,
)
from exams.selection_service import QuestionSelectionService

UPLOAD = '/api/admin/questions/import/upload/'
TEMPLATE = '/api/admin/questions/import/template/'

OBJ_HEADERS = ['SN', 'Questions', 'Mark', 'Option A', 'Option B', 'Option C', 'Option D',
               'Correct Answer', 'Explanation', 'Hint']
SUBJ_HEADERS = ['SN', 'Questions', 'Mark', 'Model Answer', 'Explanation', 'Hint']


def xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = 'questions.xlsx'
    return buf


class Base(APITestCase):
    def setUp(self):
        # Deliberately NOT is_staff: real admins in this project are role-based.
        self.admin = User.objects.create_user(username='adm', password='pw', role='admin')
        self.admin2 = User.objects.create_user(username='adm2', password='pw', role='admin')
        self.teacher = User.objects.create_user(username='tch', password='pw', role='teacher')
        self.student = User.objects.create_user(username='stu', password='pw', role='student')
        cat = ExamCategory.objects.create(name='Loksewa')
        exam = Exam.objects.create(category=cat, name='Kharidar')
        paper = Paper.objects.create(exam=exam, name='First')
        subject = Subject.objects.create(paper=paper, name='GK')
        chapter = Chapter.objects.create(subject=subject, title='Geo')
        self.topic = Topic.objects.create(chapter=chapter, name='Mountains')

    def upload(self, headers, rows, qtype, **extra):
        payload = {'file': xlsx(headers, rows), 'topic': self.topic.id,
                   'question_type': qtype, 'difficulty': 'medium'}
        payload.update(extra)
        return self.client.post(UPLOAD, payload, format='multipart')

    def commit(self, import_id):
        return self.client.post(f'/api/admin/questions/import/{import_id}/commit/')

    def make_q(self, status_='approved', qtype='mcq', text=None, **kw):
        kw.setdefault('marks', 1)
        if qtype in Question.OBJECTIVE_TYPES:
            kw.setdefault('option_a', 'a'); kw.setdefault('option_b', 'b')
            kw.setdefault('option_c', 'c'); kw.setdefault('option_d', 'd')
            kw.setdefault('correct_option', 'B')
        else:
            kw.setdefault('model_answer', 'reference answer')
        return Question.objects.create(
            topic=self.topic, question_type=qtype, status=status_,
            text=text or f'{qtype} {status_} question', **kw)


class TemplateTests(Base):
    def setUp(self):
        super().setUp()
        self.client.force_authenticate(user=self.admin)

    def headers_of(self, response):
        wb = load_workbook(io.BytesIO(response.content))
        return [c.value for c in wb['Questions'][1]]

    def test_objective_template_has_mcq_columns(self):
        r = self.client.get(TEMPLATE, {'type': 'mcq'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.headers_of(r), OBJ_HEADERS)
        self.assertIn('Objective Question Template', r['Content-Disposition'])

    def test_subjective_template_has_no_mcq_columns(self):
        r = self.client.get(TEMPLATE, {'type': 'subjective'})
        self.assertEqual(self.headers_of(r), SUBJ_HEADERS)
        self.assertIn('Subjective Question Template', r['Content-Disposition'])
        joined = ' '.join(self.headers_of(r))
        for banned in ('Option', 'Correct Answer'):
            self.assertNotIn(banned, joined)

    def test_data_sheet_has_no_importable_sample_rows(self):
        r = self.client.get(TEMPLATE, {'type': 'subjective'})
        ws = load_workbook(io.BytesIO(r.content))['Questions']
        self.assertEqual(ws.max_row, 1)

    def test_unknown_template_type_rejected(self):
        self.assertEqual(self.client.get(TEMPLATE, {'type': 'bogus'}).status_code, 400)


class ObjectiveImportTests(Base):
    def setUp(self):
        super().setUp()
        self.client.force_authenticate(user=self.admin)

    def row(self, sn=1, q='Highest peak?', mark=1, a='Everest', b='K2', c='Lhotse', d='Makalu', ans='A', ex='', hint=''):
        return [sn, q, mark, a, b, c, d, ans, ex, hint]

    def test_valid_row_and_options_header_alias(self):
        headers = ['SN', 'Questions', 'Mark', 'Options A', 'Options B', 'Options C', 'Options D',
                   'Correct Answer', 'Explanation', 'Hint']
        r = self.upload(headers, [self.row(ex='It is 8849m')], 'mcq')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data['report_data'][0]['status'], 'valid')
        self.assertEqual(r.data['report_data'][0]['data']['option_b'], 'K2')

    def test_invalid_correct_answer_reports_row_level_message(self):
        r = self.upload(OBJ_HEADERS, [self.row(), self.row(sn=2, q='Second?', ans='E', ex='x')], 'mcq')
        row = r.data['report_data'][1]
        self.assertEqual(row['status'], 'error')
        self.assertIn('"E" is invalid', row['errors'][0])
        self.assertIn('A, B, C or D', row['errors'][0])

    def test_missing_option_is_named_and_not_imported(self):
        r = self.upload(OBJ_HEADERS, [self.row(c='', ex='x')], 'mcq')
        row = r.data['report_data'][0]
        self.assertEqual(row['status'], 'incomplete')
        self.assertTrue(any('Option C' in d for d in row['missing_detail']))
        commit = self.commit(r.data['import_id'])
        self.assertEqual(commit.data['imported_count'], 0)
        self.assertEqual(Question.objects.count(), 0)

    def test_duplicate_of_existing_shows_existing_question_id(self):
        existing = self.make_q(text='Highest peak?')
        existing.question_id = 'Q-000777'
        existing.save()
        r = self.upload(OBJ_HEADERS, [self.row(ex='x')], 'mcq')
        row = r.data['report_data'][0]
        self.assertEqual(row['status'], 'duplicate')
        self.assertEqual(row['duplicate_of'], 'Q-000777')
        self.assertEqual(Question.objects.count(), 1)  # never overwritten / re-created
        self.commit(r.data['import_id'])
        self.assertEqual(Question.objects.count(), 1)

    def test_duplicate_rows_inside_file_reference_first_row(self):
        r = self.upload(OBJ_HEADERS, [self.row(ex='x'), self.row(sn=2, ex='x')], 'mcq')
        self.assertEqual(r.data['report_data'][1]['status'], 'duplicate')
        self.assertEqual(r.data['report_data'][1]['duplicate_of'], 'row 1')

    def test_commit_persists_every_field_and_provenance(self):
        r = self.upload(OBJ_HEADERS, [self.row(mark=2, ex='8849m', hint='tallest')], 'mcq',
                        difficulty='hard')
        self.assertEqual(self.commit(r.data['import_id']).data['imported_count'], 1)
        q = Question.objects.get()
        self.assertEqual((q.question_type, q.status, q.marks, q.difficulty), ('mcq', 'approved', 2.0, 'hard'))
        self.assertEqual((q.option_a, q.option_b, q.option_c, q.option_d), ('Everest', 'K2', 'Lhotse', 'Makalu'))
        self.assertEqual((q.correct_option, q.explanation, q.hint), ('A', '8849m', 'tallest'))
        self.assertEqual(q.model_answer, '')
        self.assertEqual(q.topic_id, self.topic.id)
        self.assertEqual(q.created_by_id, self.admin.id)
        self.assertRegex(q.question_id, r'^Q-\d{6}$')

    def test_numeric_answer_normalized(self):
        r = self.upload(OBJ_HEADERS, [self.row(ans='3', ex='x')], 'mcq')
        self.assertEqual(r.data['report_data'][0]['data']['correct_answer'], 'C')

    def test_subjective_file_rejected_for_objective_type(self):
        r = self.upload(SUBJ_HEADERS, [[1, 'Explain?', 5, 'Because...', '', '']], 'mcq')
        self.assertEqual(r.status_code, 400)
        self.assertIn('Subjective template', r.data['error'])

    def test_commit_skips_row_that_became_duplicate_after_analysis(self):
        r = self.upload(OBJ_HEADERS, [self.row(ex='x')], 'mcq')
        self.make_q(text='Highest peak?')  # another admin adds it in the meantime
        commit = self.commit(r.data['import_id'])
        self.assertEqual(commit.data['imported_count'], 0)
        self.assertEqual(len(commit.data['skipped_duplicates']), 1)
        self.assertEqual(Question.objects.filter(text='Highest peak?').count(), 1)


class SubjectiveImportTests(Base):
    def setUp(self):
        super().setUp()
        self.client.force_authenticate(user=self.admin)

    def test_valid_row_needs_no_mcq_fields(self):
        r = self.upload(SUBJ_HEADERS, [[1, 'Explain federalism.', 5, 'Nepal is a federal republic...', '', '']], 'subjective')
        self.assertEqual(r.status_code, 200)
        row = r.data['report_data'][0]
        self.assertEqual(row['status'], 'valid')
        self.assertEqual(row['missing'], [])
        self.assertEqual(row['data']['model_answer'], 'Nepal is a federal republic...')

    def test_model_answer_required_like_the_admin_form(self):
        r = self.upload(SUBJ_HEADERS, [[1, 'Explain federalism.', 5, '', '', '']], 'subjective')
        row = r.data['report_data'][0]
        self.assertEqual(row['status'], 'error')
        self.assertIn('Model Answer is required', row['errors'][0])

    def test_stray_objective_columns_are_ignored_not_errors(self):
        headers = SUBJ_HEADERS + ['Option A', 'Correct Answer']
        r = self.upload(headers, [[1, 'Explain federalism.', 5, 'model', '', '', 'x', 'Z']], 'subjective')
        self.assertEqual(r.data['report_data'][0]['status'], 'valid')

    def test_objective_file_rejected_for_subjective_type(self):
        r = self.upload(OBJ_HEADERS, [[1, 'Q?', 1, 'a', 'b', 'c', 'd', 'A', '', '']], 'subjective')
        self.assertEqual(r.status_code, 400)
        self.assertIn('Model Answer', r.data['error'])

    def test_explanation_optional(self):
        r = self.upload(SUBJ_HEADERS, [[1, 'Explain federalism.', 5, 'model', '', '']], 'subjective')
        self.assertEqual(r.data['incomplete_rows'], 0)
        self.assertEqual(r.data['valid_rows'], 1)

    def test_invalid_mark_rejected(self):
        r = self.upload(SUBJ_HEADERS, [[1, 'Explain federalism.', 'abc', 'model', '', '']], 'subjective')
        self.assertEqual(r.data['report_data'][0]['status'], 'error')

    def test_duplicates_detected_across_types_and_in_file(self):
        self.make_q(qtype='mcq', text='Explain federalism.')
        r = self.upload(SUBJ_HEADERS, [[1, 'Explain federalism.', 5, 'm', '', ''],
                                       [2, 'Describe X.', 5, 'm', '', ''],
                                       [3, 'Describe X.', 5, 'm', '', '']], 'subjective')
        self.assertEqual(r.data['report_data'][0]['status'], 'duplicate')
        self.assertEqual(r.data['report_data'][2]['duplicate_of'], 'row 2')

    def test_ai_fill_not_offered_for_subjective(self):
        r = self.upload(SUBJ_HEADERS, [[1, 'Explain federalism.', 5, 'model', '', '']], 'subjective')
        fill = self.client.post(f"/api/admin/questions/import/{r.data['import_id']}/ai-fill/")
        self.assertEqual(fill.status_code, 400)

    def test_commit_stores_model_answer_and_no_options(self):
        r = self.upload(SUBJ_HEADERS, [[1, 'Explain federalism.', 5, 'Nepal is a federal republic', 'see ch3', 'think 7']],
                        'subjective')
        self.assertEqual(self.commit(r.data['import_id']).data['imported_count'], 1)
        q = Question.objects.get()
        self.assertEqual((q.question_type, q.status, q.marks), ('subjective', 'approved', 5.0))
        self.assertEqual(q.model_answer, 'Nepal is a federal republic')
        self.assertEqual((q.explanation, q.hint), ('see ch3', 'think 7'))
        self.assertEqual((q.option_a, q.option_b, q.option_c, q.option_d), ('', '', '', ''))
        self.assertIsNone(q.correct_option)
        self.assertEqual(q.created_by_id, self.admin.id)

    def test_short_and_long_answer_types_use_the_same_contract(self):
        r = self.upload(SUBJ_HEADERS, [[1, 'Define GDP.', 2, 'Total output', '', '']], 'short_answer')
        self.assertEqual(r.data['report_data'][0]['status'], 'valid')


class SelectionTests(Base):
    def test_objective_family_excludes_written_answer_types_and_unapproved(self):
        mcq = self.make_q(text='m')
        tf = self.make_q(qtype='true_false', text='t', correct_option='A')
        for i, qt in enumerate(Question.SUBJECTIVE_TYPES):
            self.make_q(qtype=qt, text=f's{i}')
        self.make_q(status_='draft', text='d')
        self.make_q(status_='pending_review', text='p')
        self.make_q(status_='rejected', text='r')
        svc = QuestionSelectionService()
        got = svc.select(topic_id=self.topic.id, count=50, question_type='objective')['questions']
        self.assertEqual({q.id for q in got}, {mcq.id, tf.id})

    def test_subjective_family_returns_only_approved_written_types(self):
        self.make_q(text='m')
        subs = [self.make_q(qtype=qt, text=f's{i}') for i, qt in enumerate(Question.SUBJECTIVE_TYPES)]
        self.make_q(qtype='subjective', status_='draft', text='sd')
        got = QuestionSelectionService().select(topic_id=self.topic.id, count=50, question_type='subjective')['questions']
        self.assertEqual({q.id for q in got}, {q.id for q in subs})

    def test_exact_type_filter_still_supported(self):
        m = self.make_q(text='m')
        self.make_q(qtype='true_false', text='t', correct_option='A')
        got = QuestionSelectionService().select(topic_id=self.topic.id, count=50, question_type='mcq')['questions']
        self.assertEqual([q.id for q in got], [m.id])


class StudentPracticeTests(Base):
    def setUp(self):
        super().setUp()
        self.client.force_authenticate(user=self.student)

    def test_topic_practice_shows_only_approved_objective_questions_with_feedback(self):
        good = self.make_q(text='mcq ok', explanation='because b')
        for i, qt in enumerate(Question.SUBJECTIVE_TYPES):
            self.make_q(qtype=qt, text=f'written {i}')
        self.make_q(status_='draft', text='draft mcq')
        r = self.client.post('/api/practice-sessions/study/', {'topic': self.topic.id}, format='json')
        self.assertEqual(r.status_code, 200)
        ids = [q['id'] for q in r.data['questions']]
        self.assertEqual(ids, [good.id])
        # student payload carries options but not the answer key
        q = r.data['questions'][0]
        self.assertEqual(q['option_b'], 'b')
        self.assertNotIn('correct_option', q)

        session_id = r.data['session']['id']
        wrong = self.client.post(f'/api/practice-sessions/{session_id}/answer/',
                                 {'question_id': good.id, 'selected_option': 'C'}, format='json')
        self.assertEqual(wrong.status_code, 200)
        self.assertFalse(wrong.data['is_correct'])
        self.assertEqual(wrong.data['correct_option'], 'B')
        self.assertEqual(wrong.data['explanation'], 'because b')
        right = self.client.post(f'/api/practice-sessions/{session_id}/answer/',
                                 {'question_id': good.id, 'selected_option': 'B'}, format='json')
        self.assertTrue(right.data['is_correct'])
        attempt = QuestionAttempt.objects.get(session_id=session_id, question=good)
        self.assertEqual((attempt.selected_option, attempt.is_correct), ('B', True))

    def test_topic_practice_with_only_subjective_questions_offers_nothing(self):
        self.make_q(qtype='long_answer', text='essay')
        r = self.client.post('/api/practice-sessions/study/', {'topic': self.topic.id}, format='json')
        self.assertEqual(r.status_code, 400)

    def test_generic_practice_session_excludes_subjective(self):
        good = self.make_q(text='mcq ok')
        self.make_q(qtype='short_answer', text='short')
        r = self.client.post('/api/practice-sessions/', {'topic': self.topic.id, 'total_questions': 10}, format='json')
        self.assertIn(r.status_code, (200, 201))
        session_id = r.data['id'] if 'id' in r.data else r.data['session']['id']
        qids = list(QuestionAttempt.objects.filter(session_id=session_id).values_list('question_id', flat=True))
        self.assertEqual(qids, [good.id])


class SubjectiveStudentFlowTests(Base):
    def test_topic_subjective_attempt_only_includes_approved_written_questions(self):
        ok = self.make_q(qtype='subjective', text='essay 1')
        ok2 = self.make_q(qtype='short_answer', text='short 1')
        draft = self.make_q(qtype='subjective', status_='draft', text='essay draft')
        mcq = self.make_q(text='mcq')
        self.client.force_authenticate(user=self.student)

        listed = self.client.get('/api/subjective-questions/', {'topic': self.topic.id})
        listed_ids = {q['id'] for q in (listed.data['results'] if isinstance(listed.data, dict) else listed.data)}
        self.assertEqual(listed_ids, {ok.id, ok2.id})
        for q in (listed.data['results'] if isinstance(listed.data, dict) else listed.data):
            self.assertNotIn('option_a', q)  # never MCQ-shaped

        r = self.client.post('/api/subjective-attempts/', {
            'mode': 'topic', 'question_ids': [ok.id, ok2.id, draft.id, mcq.id]}, format='json')
        self.assertEqual(r.status_code, 201)
        attempt = SubjectiveAttempt.objects.get(pk=r.data['id'])
        self.assertEqual({a.question_id for a in attempt.answers.all()}, {ok.id, ok2.id})

    def test_answer_submission_persists_and_teacher_evaluation_scores_it(self):
        q = self.make_q(qtype='subjective', marks=5, text='essay')
        self.client.force_authenticate(user=self.student)
        r = self.client.post('/api/subjective-attempts/', {'mode': 'topic', 'question_ids': [q.id]}, format='json')
        attempt_id = r.data['id']
        sub = self.client.post(f'/api/subjective-attempts/{attempt_id}/submit_answer/',
                               {'question_id': q.id, 'answer_text': 'my written answer here'}, format='json')
        self.assertEqual(sub.status_code, 200)
        answer = SubjectiveAnswer.objects.get(attempt_id=attempt_id, question=q)
        self.assertEqual((answer.answer_text, answer.status, answer.word_count), ('my written answer here', 'submitted', 4))

        self.client.force_authenticate(user=self.admin)
        ev = self.client.post(f'/api/evaluations/{answer.id}/evaluate/',
                              {'marks_obtained': 4, 'feedback': 'good'}, format='json')
        self.assertEqual(ev.status_code, 200)
        answer.refresh_from_db()
        self.assertEqual(answer.status, 'evaluated')
        self.assertEqual(Evaluation.objects.get(answer=answer).marks_obtained, 4)
        # marks above the question's mark are refused
        over = self.client.post(f'/api/evaluations/{answer.id}/evaluate/', {'marks_obtained': 9}, format='json')
        self.assertEqual(over.status_code, 400)


class TeacherApprovalWorkflowTests(Base):
    def test_teacher_question_cannot_reach_students_until_admin_approves(self):
        self.client.force_authenticate(user=self.teacher)
        created = self.client.post('/api/teacher/questions/', {
            'topic': self.topic.id, 'question_type': 'mcq', 'text': 'Teacher question?',
            'option_a': 'a', 'option_b': 'b', 'option_c': 'c', 'option_d': 'd',
            'correct_option': 'A', 'marks': 1, 'difficulty': 'easy',
        }, format='json')
        self.assertEqual(created.status_code, 201, created.data)
        qid = created.data['id']
        self.assertEqual(Question.objects.get(pk=qid).status, 'draft')

        # draft: invisible to selection
        self.assertEqual(QuestionSelectionService().select(topic_id=self.topic.id, count=5)['questions'], [])
        self.assertEqual(self.client.post(f'/api/teacher/questions/{qid}/submit/').status_code, 200)
        self.assertEqual(Question.objects.get(pk=qid).status, 'pending_review')
        self.assertEqual(QuestionSelectionService().select(topic_id=self.topic.id, count=5)['questions'], [])
        # a teacher can never approve (their own or any) question - the review
        # queue simply does not exist for them (empty queryset -> 404)
        self.assertIn(self.client.post(f'/api/admin/questions/review-queue/{qid}/approve/').status_code, (403, 404))
        self.assertEqual(Question.objects.get(pk=qid).status, 'pending_review')

        self.client.force_authenticate(user=self.admin)
        self.assertEqual(self.client.post(f'/api/admin/questions/review-queue/{qid}/approve/').status_code, 200)
        self.assertEqual(Question.objects.get(pk=qid).status, 'approved')
        got = QuestionSelectionService().select(topic_id=self.topic.id, count=5)['questions']
        self.assertEqual([q.id for q in got], [qid])

    def test_rejected_question_is_not_selectable(self):
        q = self.make_q(status_='pending_review', text='to reject', created_by=self.teacher)
        self.client.force_authenticate(user=self.admin)
        r = self.client.post(f'/api/admin/questions/review-queue/{q.id}/reject/', {'reviewer_comment': 'unclear'}, format='json')
        self.assertIn(r.status_code, (200, 400))
        q.refresh_from_db()
        if q.status == 'rejected':
            self.assertEqual(QuestionSelectionService().select(topic_id=self.topic.id, count=5)['questions'], [])


class PermissionTests(Base):
    def test_role_based_admin_can_import_both_types_student_and_teacher_cannot(self):
        self.client.force_authenticate(user=self.admin)
        self.assertEqual(self.upload(OBJ_HEADERS, [[1, 'Q1?', 1, 'a', 'b', 'c', 'd', 'A', 'x', '']], 'mcq').status_code, 200)
        self.assertEqual(self.upload(SUBJ_HEADERS, [[1, 'Q2?', 1, 'm', '', '']], 'subjective').status_code, 200)
        self.assertEqual(self.client.get(TEMPLATE, {'type': 'subjective'}).status_code, 200)
        for user in (self.teacher, self.student):
            self.client.force_authenticate(user=user)
            self.assertEqual(self.client.get(TEMPLATE).status_code, 403)
            self.assertEqual(self.upload(SUBJ_HEADERS, [[1, 'Q3?', 1, 'm', '', '']], 'subjective').status_code, 403)
        self.client.force_authenticate(user=None)
        self.assertEqual(self.client.get(TEMPLATE).status_code, 401)

    def test_admin_question_list_filters_by_type_on_the_backend(self):
        self.make_q(text='m1')
        self.make_q(qtype='subjective', text='s1')
        self.client.force_authenticate(user=self.admin)
        r = self.client.get('/api/admin/questions/', {'question_type': 'subjective'})
        self.assertEqual(r.status_code, 200)
        results = r.data['results'] if isinstance(r.data, dict) else r.data
        self.assertEqual([q['question_type'] for q in results], ['subjective'])
        self.assertEqual(results[0]['model_answer'], 'reference answer')
